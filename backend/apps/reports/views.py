import csv
import io
from datetime import datetime
from django.db.models import Count, Sum, Avg, Q
from django.http import HttpResponse
from rest_framework import permissions, status
from rest_framework.views import APIView
from rest_framework.response import Response
from apps.audit.models import AuditEvent
from apps.locations.models import Woreda
from apps.traders.models import Trader, LegalTrader, InformalTrader, WorkflowStatus, TraderType
from apps.formalization.models import FormalizationRecord

def sanitize_csv_val(val):
    """Mitigates spreadsheet formula injection vulnerability."""
    if val is None:
        return ''
    s = str(val).strip()
    if s.startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + s
    return s

class DashboardMetricsView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        total = Trader.objects.count()
        legal = Trader.objects.filter(trader_type='LEGAL').count()
        informal = Trader.objects.filter(trader_type='INFORMAL').count()

        approved = Trader.objects.filter(status=WorkflowStatus.APPROVED).count()
        pending = Trader.objects.filter(status__in=[WorkflowStatus.SUBMITTED, 'PENDING']).count()
        under_review = Trader.objects.filter(status=WorkflowStatus.UNDER_REVIEW).count()
        needs_correction = Trader.objects.filter(status=WorkflowStatus.NEEDS_CORRECTION).count()
        rejected = Trader.objects.filter(status=WorkflowStatus.REJECTED).count()

        # Woreda breakdown
        woredas = Woreda.objects.annotate(
            total_traders=Count('traders'),
            legal_count=Count('traders', filter=Q(traders__trader_type='LEGAL')),
            informal_count=Count('traders', filter=Q(traders__trader_type='INFORMAL')),
            approved_count=Count('traders', filter=Q(traders__status=WorkflowStatus.APPROVED))
        ).order_by('id')

        woreda_stats = [
            {
                'id': w.id,
                'name': w.name,
                'code': w.code,
                'total': w.total_traders,
                'legal': w.legal_count,
                'informal': w.informal_count,
                'approved': w.approved_count
            }
            for w in woredas
        ]

        # Sector breakdown for legal
        sectors = LegalTrader.objects.values('business_sector').annotate(count=Count('id')).order_by('-count')

        # Capital assets summary for informal
        capital_agg = InformalTrader.objects.aggregate(
            total_capital=Sum('estimated_capital_assets'),
            avg_capital=Avg('estimated_capital_assets')
        )

        # Formalization breakdown
        formalization_stats = FormalizationRecord.objects.values('status').annotate(count=Count('id'))

        return Response({
            'success': True,
            'metrics': {
                'total': total,
                'legal': legal,
                'informal': informal,
                'approved': approved,
                'pending': pending,
                'underReview': under_review,
                'needsCorrection': needs_correction,
                'rejected': rejected,
                'woredaBreakdown': woreda_stats,
                'sectorBreakdown': list(sectors),
                'totalInformalCapital': float(capital_agg['total_capital'] or 0),
                'avgInformalCapital': float(capital_agg['avg_capital'] or 0),
                'formalizationStats': list(formalization_stats)
            }
        })


class CsvExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user
        qs = Trader.objects.select_related('woreda', 'kebele', 'legal_details', 'informal_details').order_by('-id')

        # Filter by woreda if requested
        woreda_id = request.query_params.get('woredaId')
        if woreda_id:
            qs = qs.filter(woreda_id=woreda_id)
        trader_type = request.query_params.get('traderType')
        if trader_type:
            qs = qs.filter(trader_type=trader_type)

        response = HttpResponse(content_type='text/csv; charset=utf-8')
        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        response['Content-Disposition'] = f'attachment; filename="HT_TMIS_Traders_{timestamp_str}.csv"'

        # BOM for UTF-8 Excel compatibility
        response.write('\ufeff'.encode('utf8'))

        writer = csv.writer(response)
        writer.writerow([
            'Trader ID', 'Type', 'Status', 'Business / Trade Name', 'Owner Full Name',
            'Phone', 'Woreda', 'Kebele', 'TIN / Resident ID', 'License / Reg Number',
            'Sector / Nature', 'Registration Date'
        ])

        for t in qs:
            is_legal = t.trader_type == 'LEGAL'
            tin_id = t.legal_details.tin if (is_legal and hasattr(t, 'legal_details')) else (
                t.informal_details.national_id_resident_id if (not is_legal and hasattr(t, 'informal_details')) else ''
            )
            reg_num = t.legal_details.trade_registration_number if (is_legal and hasattr(t, 'legal_details')) else ''
            sector = t.legal_details.business_sector if (is_legal and hasattr(t, 'legal_details')) else (
                t.informal_details.nature_of_trade_activity if (not is_legal and hasattr(t, 'informal_details')) else ''
            )

            writer.writerow([
                sanitize_csv_val(t.trader_id),
                sanitize_csv_val(t.trader_type),
                sanitize_csv_val(t.status),
                sanitize_csv_val(t.business_name),
                sanitize_csv_val(t.owner_full_name),
                sanitize_csv_val(t.phone_number),
                sanitize_csv_val(t.woreda.name if t.woreda else ''),
                sanitize_csv_val(t.kebele.name if t.kebele else ''),
                sanitize_csv_val(tin_id),
                sanitize_csv_val(reg_num),
                sanitize_csv_val(sector),
                sanitize_csv_val(t.created_at.strftime('%Y-%m-%d %H:%M')),
            ])

        AuditEvent.log(
            actor=user,
            action='EXPORT_CSV',
            entity_type='REPORT',
            entity_id=f"CSV_{timestamp_str}",
            details=f"Exported {qs.count()} trader records to CSV"
        )
        return response


class ExcelExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        user = request.user
        qs = Trader.objects.select_related('woreda', 'kebele', 'legal_details', 'informal_details').order_by('-id')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Traders Registry"

        # Styling
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

        ws.merge_cells('A1:K1')
        ws['A1'] = "Harari People Regional State — Trade & Industry Development Agency"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal="center")

        ws.merge_cells('A2:K2')
        ws['A2'] = f"HT-TMIS Official Trader Registry • Export Date: {datetime.now():%Y-%m-%d %H:%M}"
        ws['A2'].alignment = Alignment(horizontal="center")

        headers = [
            'Trader ID', 'Type', 'Status', 'Trade / Business Name', 'Owner Full Name',
            'Phone', 'Woreda', 'Kebele', 'TIN / ID', 'License No', 'Registration Date'
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for t in qs:
            row_idx += 1
            is_legal = t.trader_type == 'LEGAL'
            tin_id = t.legal_details.tin if (is_legal and hasattr(t, 'legal_details')) else (
                t.informal_details.national_id_resident_id if (not is_legal and hasattr(t, 'informal_details')) else ''
            )
            reg_num = t.legal_details.trade_registration_number if (is_legal and hasattr(t, 'legal_details')) else ''

            ws.append([
                sanitize_csv_val(t.trader_id),
                t.trader_type,
                t.status,
                sanitize_csv_val(t.business_name),
                sanitize_csv_val(t.owner_full_name),
                sanitize_csv_val(t.phone_number),
                t.woreda.name if t.woreda else '',
                t.kebele.name if t.kebele else '',
                sanitize_csv_val(tin_id),
                sanitize_csv_val(reg_num),
                t.created_at.strftime('%Y-%m-%d %H:%M')
            ])

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="HT_TMIS_Registry_{timestamp_str}.xlsx"'

        AuditEvent.log(
            actor=user,
            action='EXPORT_EXCEL',
            entity_type='REPORT',
            entity_id=f"EXCEL_{timestamp_str}",
            details=f"Exported {qs.count()} trader records to Excel"
        )
        return response


class PdfExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        user = request.user
        qs = Trader.objects.select_related('woreda', 'kebele').order_by('-id')[:50]

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#1E3A8A'),
            alignment=1,
            spaceAfter=6
        )
        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#4B5563'),
            alignment=1,
            spaceAfter=14
        )

        elements.append(Paragraph("HARARI PEOPLE REGIONAL STATE", title_style))
        elements.append(Paragraph("Trade & Industry Development Agency — Official Trader Registry", title_style))
        elements.append(Paragraph(f"Generated on {datetime.now():%B %d, %Y at %H:%M} • Authorized Officer: {user.full_name}", sub_style))

        table_data = [
            ['Trader ID', 'Type', 'Status', 'Trade / Owner Name', 'Woreda', 'Kebele', 'Date Registered']
        ]

        for t in qs:
            table_data.append([
                t.trader_id,
                t.trader_type,
                t.status,
                t.owner_full_name[:25],
                t.woreda.name if t.woreda else '',
                t.kebele.name if t.kebele else '',
                t.created_at.strftime('%Y-%m-%d')
            ])

        t = Table(table_data, colWidths=[90, 65, 85, 200, 100, 100, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
        ]))

        elements.append(t)
        doc.build(elements)
        buffer.seek(0)

        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="HT_TMIS_Summary_{timestamp_str}.pdf"'

        AuditEvent.log(
            actor=user,
            action='EXPORT_PDF',
            entity_type='REPORT',
            entity_id=f"PDF_{timestamp_str}",
            details=f"Exported official PDF report"
        )
        return response
