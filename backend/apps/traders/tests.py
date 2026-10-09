import io
from decimal import Decimal
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient
from rest_framework import status

from apps.accounts.models import User
from apps.locations.models import Region, Woreda, Kebele
from apps.traders.models import Trader, LegalTrader, InformalTrader
from apps.formalization.models import FormalizationAssessment
from apps.audit.models import AuditLog


class HTTMISIntegrationTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Create Region, Woredas, and Kebeles
        self.region = Region.objects.create(name='Harari Region', code='HR-01')
        self.woreda_an = Woreda.objects.create(name='Amir Nur', code='AN-01', region=self.region)
        self.kebele_an1 = Kebele.objects.create(name='Kebele 01 (Jenila)', code='AN-K01', woreda=self.woreda_an)

        self.woreda_ab = Woreda.objects.create(name='Abadir', code='AB-02', region=self.region)
        self.kebele_ab4 = Kebele.objects.create(name='Kebele 04 (Gedir Shingir)', code='AB-K04', woreda=self.woreda_ab)

        # Users
        self.encoder = User.objects.create_user(
            username='encoder1',
            password='password123',
            full_name='Abebe Data Encoder',
            role='DATA_ENCODER'
        )
        self.director = User.objects.create_user(
            username='director1',
            password='password123',
            full_name='Dr. Kebede Director',
            role='DIRECTOR'
        )
        self.admin = User.objects.create_user(
            username='admin1',
            password='password123',
            full_name='System Admin',
            role='SYSTEM_ADMINISTRATOR',
            is_staff=True
        )

    def test_woreda_kebele_mismatch_rejected(self):
        """Woreda and Kebele relationships must be validated server-side."""
        self.client.force_authenticate(user=self.encoder)

        # Mismatch: Kebele from Abadir passed with Amir Nur woreda ID
        payload = {
            'tradeName': 'Invalid Location Cafe',
            'ownerFullName': 'Test Owner',
            'phoneNumber': '+251911000000',
            'tin': '9988776655',
            'tradeRegistrationNumber': 'HR-TR-TEST-001',
            'woredaId': self.woreda_an.id,
            'kebeleId': self.kebele_ab4.id, # Kebele from Abadir!
            'gender': 'MALE',
            'age': 30,
            'businessSector': 'GENERAL_TRADE',
            'tradeScale': 'RETAIL',
            'businessOwnershipType': 'SOLE_PROPRIETORSHIP',
            'issuingInstitution': 'Harari Trade Bureau',
            'dateOfIssuance': '2024-01-01',
        }
        res = self.client.post('/api/v1/traders/legal/', payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.data.get('code'), 'INVALID_WOREDA_KEBELE_COMBINATION')

    def test_legal_trader_registration_and_duplicate_tin_rejection(self):
        """Register legal trader, then verify duplicate TIN is rejected."""
        self.client.force_authenticate(user=self.encoder)

        payload = {
            'tradeName': 'Harar Spice PLC',
            'ownerFullName': 'Dawit Spice',
            'phoneNumber': '+251911223344',
            'tin': '1122334455',
            'tradeRegistrationNumber': 'HR-TR-2026-SPICE',
            'woredaId': self.woreda_an.id,
            'kebeleId': self.kebele_an1.id,
            'gender': 'MALE',
            'age': 40,
            'businessSector': 'AGRICULTURE_AGRO_PROCESSING',
            'tradeScale': 'WHOLESALE',
            'businessOwnershipType': 'PLC',
            'issuingInstitution': 'Harari Trade Bureau',
            'dateOfIssuance': '2025-05-10',
        }
        res = self.client.post('/api/v1/traders/legal/', payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['traderType'], 'LEGAL')
        trader_id = res.data['traderId']

        # Attempt duplicate TIN
        res_dup = self.client.post('/api/v1/traders/legal/', payload, format='json')
        self.assertEqual(res_dup.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res_dup.data.get('code'), 'DUPLICATE_TIN')

    def test_informal_trader_registration_and_formalization(self):
        """Register informal trader and verify formalization remains separate."""
        self.client.force_authenticate(user=self.encoder)

        payload = {
            'fullName': 'Chaltu Gemechu',
            'phoneNumber': '+251933445566',
            'nationalIdResidentId': 'ET-HR-CHALTU-01',
            'woredaId': self.woreda_ab.id,
            'kebeleId': self.kebele_ab4.id,
            'specificLocationMarketArea': 'Gedir Shingir Market',
            'gender': 'FEMALE',
            'age': 29,
            'natureOfTradeActivity': 'STREET_VENDING_OPEN_MARKET',
            'estimatedCapitalAssets': 15000.0,
            'reasonForOperatingInformally': 'LACK_OF_CAPITAL',
            'formalizationStatusRecommendation': 'READY_FOR_TIN_MICRO_ENTERPRISE',
            'enumeratorDataCollectorName': 'Abebe Data Encoder',
            'dateOfAssessment': '2026-10-08',
        }
        res = self.client.post('/api/v1/traders/informal/', payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['traderType'], 'INFORMAL')
        trader_id = res.data['traderId']

        # Director approves registration
        self.client.force_authenticate(user=self.director)
        dec_res = self.client.post(
            f'/api/v1/verification/{trader_id}/decision/',
            {'status': 'APPROVED', 'notes': 'Field assessment confirmed'},
            format='json'
        )
        self.assertEqual(dec_res.status_code, status.HTTP_200_OK)

        # Ensure approved informal trader is STILL an informal trader (not converted to legal)
        trader = Trader.objects.get(trader_id=trader_id)
        self.assertEqual(trader.status, 'APPROVED')
        self.assertEqual(trader.trader_type, 'INFORMAL')

        # Assess formalization separately
        form_res = self.client.post(
            '/api/v1/formalization/assess/',
            {'trader_id': trader_id, 'status': 'READY_FOR_FORMALIZATION', 'support_package': 'Market stall grant'},
            format='json'
        )
        self.assertEqual(form_res.status_code, status.HTTP_200_OK)
        self.assertEqual(form_res.data['status'], 'READY_FOR_FORMALIZATION')

    def test_director_cannot_approve_own_submission(self):
        """Conflict of interest guard: Director cannot approve record they submitted."""
        # Director creates a submission
        trader = Trader.objects.create(
            trader_id='HTT-SELF-01',
            trader_type='LEGAL',
            status='SUBMITTED',
            name='Director Own Venture',
            owner_full_name='Dr. Kebede Director',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
            created_by=self.director,
        )
        LegalTrader.objects.create(
            trader=trader,
            tin='9999999999',
            trade_registration_number='HR-DIR-001',
            date_of_issuance='2025-01-01',
        )

        self.client.force_authenticate(user=self.director)
        res = self.client.post(
            f'/api/v1/verification/{trader.trader_id}/decision/',
            {'status': 'APPROVED', 'notes': 'Self approving'},
            format='json'
        )
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
        self.assertIn('Conflict of interest', res.data.get('detail', ''))

    def test_admin_cannot_approve_trader(self):
        """Administrator cannot approve operational records without explicit Director role."""
        trader = Trader.objects.create(
            trader_id='HTT-TEST-ADMIN-01',
            trader_type='LEGAL',
            status='SUBMITTED',
            name='Test Trader',
            owner_full_name='Test Owner',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
            created_by=self.encoder,
        )

        self.client.force_authenticate(user=self.admin)
        res = self.client.post(
            f'/api/v1/verification/{trader.trader_id}/decision/',
            {'status': 'APPROVED', 'notes': 'Admin trying to approve'},
            format='json'
        )
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_rejection_requires_mandatory_reason(self):
        """Rejection or return requires a mandatory explanatory reason."""
        trader = Trader.objects.create(
            trader_id='HTT-TEST-REJ-01',
            trader_type='LEGAL',
            status='SUBMITTED',
            name='Test Trader Rej',
            owner_full_name='Test Owner',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
            created_by=self.encoder,
        )

        self.client.force_authenticate(user=self.director)
        # Attempt rejection without reason
        res = self.client.post(
            f'/api/v1/verification/{trader.trader_id}/decision/',
            {'status': 'REJECTED', 'notes': ''},
            format='json'
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_multilingual_csv_export(self):
        """CSV export supports en, om, am with UTF-8 BOM."""
        self.client.force_authenticate(user=self.encoder)

        for lang in ['en', 'om', 'am']:
            res = self.client.get(f'/api/v1/reports/export/csv/?lang={lang}')
            self.assertEqual(res.status_code, status.HTTP_200_OK)
            self.assertTrue(res.content.startswith(b'\xef\xbb\xbf')) # UTF-8 BOM

    def test_excel_export(self):
        """Excel export generates valid XLSX file."""
        self.client.force_authenticate(user=self.encoder)
        res = self.client.get('/api/v1/reports/export/excel/?lang=en')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    def test_pdf_certificate_generation_and_unapproved_block(self):
        """Approved legal trader can generate PDF certificate; unapproved cannot."""
        # Unapproved
        trader_unapp = Trader.objects.create(
            trader_id='HTT-UNAPP-01',
            trader_type='LEGAL',
            status='SUBMITTED',
            name='Unapproved Co',
            owner_full_name='Unapproved Owner',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )
        LegalTrader.objects.create(
            trader=trader_unapp,
            tin='1234567890',
            trade_registration_number='HR-UNAPP-01',
            date_of_issuance='2025-01-01',
        )

        self.client.force_authenticate(user=self.director)
        res_fail = self.client.get(f'/api/v1/reports/certificate/{trader_unapp.trader_id}/pdf/')
        self.assertEqual(res_fail.status_code, status.HTTP_400_BAD_REQUEST)

        # Approved
        trader_unapp.status = 'APPROVED'
        trader_unapp.save()

        res_ok = self.client.get(f'/api/v1/reports/certificate/{trader_unapp.trader_id}/pdf/?lang=am')
        self.assertEqual(res_ok.status_code, status.HTTP_200_OK)
        self.assertEqual(res_ok['Content-Type'], 'application/pdf')
        self.assertTrue(res_ok.content.startswith(b'%PDF'))

    def test_agency_leader_cannot_approve_trader(self):
        """Agency Leader cannot approve operational records without explicit Director role."""
        leader = User.objects.create_user(
            username='leader1',
            password='password123',
            full_name='Agency Leader One',
            role='AGENCY_LEADER'
        )
        trader = Trader.objects.create(
            trader_id='HTT-TEST-LEAD-01',
            trader_type='LEGAL',
            status='SUBMITTED',
            name='Test Trader Leader',
            owner_full_name='Test Owner',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
            created_by=self.encoder,
        )

        self.client.force_authenticate(user=leader)
        res = self.client.post(
            f'/api/v1/verification/{trader.trader_id}/decision/',
            {'status': 'APPROVED', 'notes': 'Leader trying to approve'},
            format='json'
        )
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_search_across_fields(self):
        """Search matches trader name, trader code, TIN, and phone number."""
        self.client.force_authenticate(user=self.director)
        trader = Trader.objects.create(
            trader_id='HTT-SRCH-999',
            trader_type='LEGAL',
            status='APPROVED',
            name='Awash Coffee Roastery',
            owner_full_name='Dawit Roaster',
            phone_number='+251911998877',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )
        LegalTrader.objects.create(
            trader=trader,
            tin='8877665544',
            trade_registration_number='HR-SRCH-999',
            date_of_issuance='2025-01-01',
        )

        # 1. Search by name
        res = self.client.get('/api/v1/traders/?search=Awash')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data['results'] if 'results' in res.data else res.data
        self.assertTrue(any(t['traderId'] == 'HTT-SRCH-999' for t in results))

        # 2. Search by trader ID
        res = self.client.get('/api/v1/traders/?search=SRCH-999')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data['results'] if 'results' in res.data else res.data
        self.assertTrue(any(t['traderId'] == 'HTT-SRCH-999' for t in results))

        # 3. Search by TIN
        res = self.client.get('/api/v1/traders/?search=8877665544')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data['results'] if 'results' in res.data else res.data
        self.assertTrue(any(t['traderId'] == 'HTT-SRCH-999' for t in results))

        # 4. Search by Phone
        res = self.client.get('/api/v1/traders/?search=911998877')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data['results'] if 'results' in res.data else res.data
        self.assertTrue(any(t['traderId'] == 'HTT-SRCH-999' for t in results))

    def test_combined_filters_and_reset(self):
        """Combining multiple filters narrows results; resetting clears them."""
        self.client.force_authenticate(user=self.director)
        t_legal = Trader.objects.create(
            trader_id='HTT-COMB-01',
            trader_type='LEGAL',
            status='APPROVED',
            name='Legal Approved AmirNur',
            owner_full_name='Owner 1',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )
        LegalTrader.objects.create(
            trader=t_legal,
            tin='1111000011',
            trade_registration_number='HR-COMB-01',
            business_sector='AGRICULTURE_AGRO_PROCESSING',
            date_of_issuance='2025-01-01',
        )

        t_informal = Trader.objects.create(
            trader_id='HTT-COMB-02',
            trader_type='INFORMAL',
            status='SUBMITTED',
            name='Informal Pending Abadir',
            owner_full_name='Owner 2',
            woreda=self.woreda_ab,
            kebele=self.kebele_ab4,
        )
        InformalTrader.objects.create(
            trader=t_informal,
            estimated_capital_assets=Decimal('5000'),
            date_of_assessment='2025-01-01',
        )

        # Combined filter matching t_legal
        res = self.client.get(f'/api/v1/traders/?trader_type=LEGAL&status=APPROVED&woreda={self.woreda_an.id}')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        results = res.data['results'] if 'results' in res.data else res.data
        self.assertTrue(any(t['traderId'] == 'HTT-COMB-01' for t in results))
        self.assertFalse(any(t['traderId'] == 'HTT-COMB-02' for t in results))

        # Reset (no filters) includes both
        res_reset = self.client.get('/api/v1/traders/')
        self.assertEqual(res_reset.status_code, status.HTTP_200_OK)
        results_all = res_reset.data['results'] if 'results' in res_reset.data else res_reset.data
        self.assertTrue(any(t['traderId'] == 'HTT-COMB-01' for t in results_all))
        self.assertTrue(any(t['traderId'] == 'HTT-COMB-02' for t in results_all))

    def test_server_side_pagination(self):
        """Server-side pagination returns 25 per page with next link for >25 records."""
        self.client.force_authenticate(user=self.director)
        for i in range(30):
            Trader.objects.create(
                trader_id=f'HTT-PAG-{i:03d}',
                trader_type='LEGAL',
                status='APPROVED',
                name=f'Pagination Trader {i}',
                owner_full_name=f'Owner {i}',
                woreda=self.woreda_an,
                kebele=self.kebele_an1,
            )

        res_p1 = self.client.get('/api/v1/traders/?page=1&page_size=25')
        self.assertEqual(res_p1.status_code, status.HTTP_200_OK)
        self.assertIn('results', res_p1.data)
        self.assertEqual(len(res_p1.data['results']), 25)
        self.assertIsNotNone(res_p1.data['next'])
        self.assertGreaterEqual(res_p1.data['count'], 30)

        res_p2 = self.client.get('/api/v1/traders/?page=2&page_size=25')
        self.assertEqual(res_p2.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(len(res_p2.data['results']), 5)
        self.assertIsNotNone(res_p2.data['previous'])

    def test_selected_excel_and_pdf_exports(self):
        """Export Selected exports strictly the selected trader IDs in Excel and PDF."""
        self.client.force_authenticate(user=self.director)
        t1 = Trader.objects.create(
            trader_id='HTT-SEL-001',
            trader_type='LEGAL',
            status='APPROVED',
            name='Selected Trader One',
            owner_full_name='Owner One',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )
        t2 = Trader.objects.create(
            trader_id='HTT-SEL-002',
            trader_type='INFORMAL',
            status='SUBMITTED',
            name='Selected Trader Two',
            owner_full_name='Owner Two',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )
        t3 = Trader.objects.create(
            trader_id='HTT-NOTSEL-003',
            trader_type='LEGAL',
            status='APPROVED',
            name='Unselected Trader Three',
            owner_full_name='Owner Three',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )

        # Excel Export Selected
        payload = {'ids': ['HTT-SEL-001', 'HTT-SEL-002'], 'scope': 'selected', 'lang': 'en'}
        res_xl = self.client.post('/api/v1/reports/export/excel/', payload, format='json')
        self.assertEqual(res_xl.status_code, status.HTTP_200_OK)
        self.assertEqual(res_xl['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res_xl.content))
        ws = wb.active
        # Data starts row 6
        rows = list(ws.iter_rows(values_only=True))[5:]
        exported_ids = [r[0] for r in rows if r[0]]
        self.assertIn('HTT-SEL-001', exported_ids)
        self.assertIn('HTT-SEL-002', exported_ids)
        self.assertNotIn('HTT-NOTSEL-003', exported_ids)

        # PDF Export Selected
        res_pdf = self.client.post('/api/v1/reports/export/pdf/', payload, format='json')
        self.assertEqual(res_pdf.status_code, status.HTTP_200_OK)
        self.assertEqual(res_pdf['Content-Type'], 'application/pdf')
        self.assertTrue(res_pdf.content.startswith(b'%PDF'))

    def test_empty_selection_and_empty_results_handling(self):
        """Empty selection returns 400 Bad Request; filtered search with 0 matches returns 200 with headers."""
        self.client.force_authenticate(user=self.director)

        # Empty selection returns 400
        res_empty_sel_xl = self.client.post(
            '/api/v1/reports/export/excel/',
            {'ids': [], 'scope': 'selected'},
            format='json'
        )
        self.assertEqual(res_empty_sel_xl.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res_empty_sel_xl.data.get('code'), 'NO_RECORDS_FOR_EXPORT')

        res_empty_sel_pdf = self.client.post(
            '/api/v1/reports/export/pdf/',
            {'ids': [], 'scope': 'selected'},
            format='json'
        )
        self.assertEqual(res_empty_sel_pdf.status_code, status.HTTP_400_BAD_REQUEST)

        # Filtered with non-matching search returns 200 empty report
        res_zero_filtered_xl = self.client.get('/api/v1/reports/export/excel/?search=NON_EXISTENT_STRING_99999')
        self.assertEqual(res_zero_filtered_xl.status_code, status.HTTP_200_OK)

        res_zero_filtered_pdf = self.client.get('/api/v1/reports/export/pdf/?search=NON_EXISTENT_STRING_99999')
        self.assertEqual(res_zero_filtered_pdf.status_code, status.HTTP_200_OK)

    def test_tampered_or_unauthorized_ids_export(self):
        """Tampered or invalid IDs in selection request return 400 Bad Request."""
        self.client.force_authenticate(user=self.director)
        payload = {'ids': ['HTT-FAKE-001', 'HTT-TAMPERED-999'], 'scope': 'selected'}
        res = self.client.post('/api/v1/reports/export/excel/', payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.data.get('code'), 'NO_RECORDS_FOR_EXPORT')

    def test_cross_territory_export_blocked_for_data_encoder(self):
        """Data Encoder restricted to Amir Nur cannot export records from Abadir."""
        # Restrict encoder to Amir Nur Woreda
        self.encoder.assigned_woreda = self.woreda_an
        self.encoder.save()

        # Create record in Amir Nur
        t_an = Trader.objects.create(
            trader_id='HTT-AN-01',
            trader_type='LEGAL',
            status='APPROVED',
            name='Amir Nur Trader',
            owner_full_name='Owner AN',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )

        # Create record in Abadir
        t_ab = Trader.objects.create(
            trader_id='HTT-AB-01',
            trader_type='LEGAL',
            status='APPROVED',
            name='Abadir Trader',
            owner_full_name='Owner AB',
            woreda=self.woreda_ab,
            kebele=self.kebele_ab4,
        )

        self.client.force_authenticate(user=self.encoder)

        # Attempt to export Abadir trader
        res_ab = self.client.post(
            '/api/v1/reports/export/excel/',
            {'ids': [t_ab.trader_id], 'scope': 'selected'},
            format='json'
        )
        self.assertEqual(res_ab.status_code, status.HTTP_400_BAD_REQUEST)

        # Filtered export includes only Amir Nur trader
        res_filtered = self.client.get('/api/v1/reports/export/excel/')
        self.assertEqual(res_filtered.status_code, status.HTTP_200_OK)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res_filtered.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))[5:]
        exported_ids = [r[0] for r in rows if r[0]]
        self.assertIn(t_an.trader_id, exported_ids)
        self.assertNotIn(t_ab.trader_id, exported_ids)

    def test_audit_logging_on_exports(self):
        """Export actions generate comprehensive AuditLog entries with format, scope, and user."""
        self.client.force_authenticate(user=self.director)
        trader = Trader.objects.create(
            trader_id='HTT-AUD-01',
            trader_type='LEGAL',
            status='APPROVED',
            name='Audit Test Trader',
            owner_full_name='Audit Owner',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )

        AuditLog.objects.all().delete()

        # Excel export
        res_xl = self.client.post(
            '/api/v1/reports/export/excel/',
            {'ids': [trader.trader_id], 'scope': 'selected', 'lang': 'om'},
            format='json'
        )
        self.assertEqual(res_xl.status_code, status.HTTP_200_OK)
        self.assertTrue(AuditLog.objects.filter(action='EXPORT_EXCEL', user=self.director.username).exists())
        log_xl = AuditLog.objects.filter(action='EXPORT_EXCEL').first()
        self.assertIn('Scope: Selected', log_xl.details)

        # PDF export
        res_pdf = self.client.post(
            '/api/v1/reports/export/pdf/',
            {'ids': [trader.trader_id], 'scope': 'selected', 'lang': 'am'},
            format='json'
        )
        self.assertEqual(res_pdf.status_code, status.HTTP_200_OK)
        self.assertTrue(AuditLog.objects.filter(action='EXPORT_PDF', user=self.director.username).exists())

    def test_excel_formula_injection_protection(self):
        """Formula injection dangerous prefixes are safely escaped with single quote in Excel."""
        self.client.force_authenticate(user=self.director)
        trader = Trader.objects.create(
            trader_id='HTT-INJ-01',
            trader_type='LEGAL',
            status='APPROVED',
            name="=CMD|' /C calc'!A0",
            owner_full_name='+251911000000',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )

        res = self.client.post(
            '/api/v1/reports/export/excel/',
            {'ids': [trader.trader_id], 'scope': 'selected'},
            format='json'
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # Check cell D6 (row 6, column 4 is name)
        row6 = list(ws.iter_rows(values_only=True))[5]
        name_val = row6[3]
        self.assertTrue(name_val.startswith("'="), f"Expected '= but got {name_val}")

    def test_certificate_pdf_cross_territory_scoping(self):
        """Certificate PDF endpoint scopes records by territory for Data Encoders."""
        # Restrict encoder to Amir Nur Woreda
        self.encoder.assigned_woreda = self.woreda_an
        self.encoder.save()

        # Approved legal trader in Amir Nur
        t_an = Trader.objects.create(
            trader_id='HTT-CERT-AN',
            trader_type='LEGAL',
            status='APPROVED',
            name='Amir Nur Cert Trader',
            owner_full_name='Owner AN',
            woreda=self.woreda_an,
            kebele=self.kebele_an1,
        )
        LegalTrader.objects.create(
            trader=t_an,
            tin='1122334455',
            trade_registration_number='HR-CERT-AN-01',
            date_of_issuance='2025-01-01',
        )

        # Approved legal trader in Abadir
        t_ab = Trader.objects.create(
            trader_id='HTT-CERT-AB',
            trader_type='LEGAL',
            status='APPROVED',
            name='Abadir Cert Trader',
            owner_full_name='Owner AB',
            woreda=self.woreda_ab,
            kebele=self.kebele_ab4,
        )
        LegalTrader.objects.create(
            trader=t_ab,
            tin='9988776655',
            trade_registration_number='HR-CERT-AB-01',
            date_of_issuance='2025-01-01',
        )

        self.client.force_authenticate(user=self.encoder)

        # 1. Access trader from own assigned territory -> Allowed (200 OK)
        res_own = self.client.get(f'/api/v1/reports/certificate/{t_an.trader_id}/pdf/')
        self.assertEqual(res_own.status_code, status.HTTP_200_OK)
        self.assertEqual(res_own['Content-Type'], 'application/pdf')

        # 2. Access trader from another territory -> Denied (404 Not Found)
        res_other = self.client.get(f'/api/v1/reports/certificate/{t_ab.trader_id}/pdf/')
        self.assertEqual(res_other.status_code, status.HTTP_404_NOT_FOUND)

        # 3. Regional Director can access both territories
        self.client.force_authenticate(user=self.director)
        res_dir_an = self.client.get(f'/api/v1/reports/certificate/{t_an.trader_id}/pdf/')
        res_dir_ab = self.client.get(f'/api/v1/reports/certificate/{t_ab.trader_id}/pdf/')
        self.assertEqual(res_dir_an.status_code, status.HTTP_200_OK)
        self.assertEqual(res_dir_ab.status_code, status.HTTP_200_OK)

    def test_export_endpoints_csrf_protection_on_post(self):
        """Export endpoints enforce CSRF validation on session-authenticated POST requests."""
        from django.test import Client
        from django.middleware.csrf import _get_new_csrf_string

        session_client = Client(enforce_csrf_checks=True)
        session_client.force_login(self.director)

        endpoints = [
            '/api/v1/reports/export/excel/',
            '/api/v1/reports/export/csv/',
            '/api/v1/reports/export/pdf/',
        ]

        # 1. POST requests without CSRF token must fail with 403 Forbidden
        for endpoint in endpoints:
            res_no_csrf = session_client.post(
                endpoint,
                {'ids': ['HTT-NONEXISTENT'], 'scope': 'selected'},
                content_type='application/json'
            )
            self.assertEqual(
                res_no_csrf.status_code,
                status.HTTP_403_FORBIDDEN,
                f"Expected 403 for {endpoint} without CSRF token"
            )

        # 2. POST requests with valid CSRF cookie and X-CSRFToken header pass CSRF check
        csrf_secret = _get_new_csrf_string()
        session_client.cookies['csrftoken'] = csrf_secret

        for endpoint in endpoints:
            res_with_csrf = session_client.post(
                endpoint,
                {'ids': ['HTT-NONEXISTENT'], 'scope': 'selected'},
                content_type='application/json',
                HTTP_X_CSRFTOKEN=csrf_secret
            )
            # Reaches view logic and fails with 400 (NO_RECORDS_FOR_EXPORT), NOT 403 CSRF error
            self.assertEqual(
                res_with_csrf.status_code,
                status.HTTP_400_BAD_REQUEST,
                f"Expected 400 (view validation) rather than 403 CSRF error for {endpoint}"
            )


